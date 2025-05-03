import os
import logging
import json
import tempfile
import mimetypes
from pathlib import Path
from typing import Dict, List, Any, Optional, Union

# File processing libraries
try:
    import chardet
    CHARDET_AVAILABLE = True
except ImportError:
    CHARDET_AVAILABLE = False

try:
    import docx2txt
    DOCX2TXT_AVAILABLE = True
except ImportError:
    DOCX2TXT_AVAILABLE = False

try:
    import PyPDF2
    PYPDF2_AVAILABLE = True
except ImportError:
    PYPDF2_AVAILABLE = False

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    import pptx
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    import xml.etree.ElementTree as ET
    XML_AVAILABLE = True
except ImportError:
    XML_AVAILABLE = False

try:
    import zipfile
    import tarfile
    ARCHIVE_AVAILABLE = True
except ImportError:
    ARCHIVE_AVAILABLE = False

logger = logging.getLogger(__name__)

class FileProcessor:
    """Processor for various file types."""
    
    def __init__(self, config=None):
        """
        Initialize the file processor.
        
        Args:
            config (dict, optional): Configuration for file processing
        """
        self.config = config or {}
        self.temp_dir = tempfile.mkdtemp()
        self.max_file_size = self.config.get("max_file_size_mb", 10) * 1024 * 1024  # Convert to bytes
        self.allowed_extensions = self.config.get("allowed_file_types")
    
    def process_file(self, file_path: str) -> Dict[str, Any]:
        """
        Process a file and extract its content based on file type.
        
        Args:
            file_path (str): Path to the file
            
        Returns:
            Dict[str, Any]: The extracted content and metadata
        """
        if not os.path.exists(file_path):
            return {"error": f"File {file_path} not found"}
            
        # Check file size
        file_size = os.path.getsize(file_path)
        if file_size > self.max_file_size:
            return {"error": f"File size ({file_size} bytes) exceeds maximum allowed size ({self.max_file_size} bytes)"}
            
        # Check file extension if restrictions are in place
        if self.allowed_extensions:
            file_extension = os.path.splitext(file_path)[1].lower()
            if file_extension not in self.allowed_extensions:
                return {"error": f"File type {file_extension} is not allowed"}
        
        try:
            file_extension = os.path.splitext(file_path)[1].lower()
            
            # Process based on file extension
            if file_extension in ['.py', '.js', '.html', '.css', '.json', '.md', '.txt', '.csv', '.xml', '.sh']:
                return self._process_text_file(file_path, file_extension)
                
            elif file_extension == '.pdf':
                return self._process_pdf(file_path)
                
            elif file_extension == '.docx':
                return self._process_docx(file_path)
                
            elif file_extension in ['.xlsx', '.xls']:
                return self._process_excel(file_path)
                
            elif file_extension == '.pptx':
                return self._process_powerpoint(file_path)
                
            elif file_extension in ['.zip', '.tar', '.gz']:
                return self._process_archive(file_path, file_extension)
                
            else:
                # Try to process as a text file
                return self._process_unknown_file(file_path)
                
        except Exception as e:
            logger.error(f"Error processing file {file_path}: {str(e)}", exc_info=True)
            return {"error": f"Failed to process file: {str(e)}"}
    
    def _process_text_file(self, file_path: str, file_extension: str) -> Dict[str, Any]:
        """Process a text-based file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
                
            file_type = file_extension.lstrip('.')
            if file_extension == '.json':
                try:
                    content = json.loads(content)
                except json.JSONDecodeError:
                    return {"error": "Invalid JSON file", "raw_content": content}
                    
            elif file_extension == '.csv' and PANDAS_AVAILABLE:
                try:
                    df = pd.read_csv(file_path)
                    return {
                        "type": "csv",
                        "content": content,
                        "dataframe": df.to_dict(),
                        "columns": df.columns.tolist(),
                        "shape": df.shape
                    }
                except Exception as e:
                    return {"error": f"Error parsing CSV: {str(e)}", "raw_content": content}
                    
            elif file_extension == '.xml' and XML_AVAILABLE:
                try:
                    tree = ET.parse(file_path)
                    root = tree.getroot()
                    return {
                        "type": "xml",
                        "content": content,
                        "root_tag": root.tag,
                        "structure": self._xml_to_dict(root)
                    }
                except Exception as e:
                    return {"error": f"Error parsing XML: {str(e)}", "raw_content": content}
            
            return {
                "type": file_type,
                "content": content
            }
            
        except UnicodeDecodeError:
            # Try to detect encoding
            if CHARDET_AVAILABLE:
                with open(file_path, 'rb') as f:
                    rawdata = f.read()
                result = chardet.detect(rawdata)
                encoding = result['encoding']
                
                with open(file_path, 'r', encoding=encoding) as f:
                    content = f.read()
                return {
                    "type": file_extension.lstrip('.'),
                    "content": content,
                    "encoding": encoding
                }
            else:
                return {"error": "Unable to decode file and chardet is not available"}
    
    def _process_pdf(self, file_path: str) -> Dict[str, Any]:
        """Process a PDF file."""
        if not PYPDF2_AVAILABLE:
            return {"error": "PyPDF2 is not available. Please install with: pip install PyPDF2"}
            
        try:
            text = ""
            metadata = {}
            
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                metadata = {
                    "pages": len(reader.pages),
                    "encrypted": reader.is_encrypted
                }
                
                # Extract document info if available
                if hasattr(reader, 'metadata') and reader.metadata:
                    for key, value in reader.metadata.items():
                        if key.startswith('/'):
                            key = key[1:]
                        metadata[key] = value
                
                # Extract text from each page
                for page in reader.pages:
                    text += page.extract_text() + "\n\n"
            
            return {
                "type": "pdf",
                "content": text,
                "metadata": metadata
            }
            
        except Exception as e:
            logger.error(f"Error processing PDF: {str(e)}", exc_info=True)
            return {"error": f"Failed to process PDF: {str(e)}"}
    
    def _process_docx(self, file_path: str) -> Dict[str, Any]:
        """Process a Word document."""
        if not DOCX2TXT_AVAILABLE:
            return {"error": "docx2txt is not available. Please install with: pip install docx2txt"}
            
        try:
            text = docx2txt.process(file_path)
            return {
                "type": "docx",
                "content": text
            }
            
        except Exception as e:
            logger.error(f"Error processing DOCX: {str(e)}", exc_info=True)
            return {"error": f"Failed to process DOCX: {str(e)}"}
    
    def _process_excel(self, file_path: str) -> Dict[str, Any]:
        """Process an Excel file."""
        if not PANDAS_AVAILABLE or not OPENPYXL_AVAILABLE:
            return {"error": "pandas and openpyxl are required. Please install with: pip install pandas openpyxl"}
            
        try:
            # Get sheet names
            workbook = openpyxl.load_workbook(file_path, read_only=True)
            sheet_names = workbook.sheetnames
            workbook.close()
            
            # Read all sheets
            sheets_data = {}
            excel_file = pd.ExcelFile(file_path)
            
            for sheet_name in sheet_names:
                df = pd.read_excel(excel_file, sheet_name=sheet_name)
                sheets_data[sheet_name] = {
                    "data": df.to_dict(),
                    "columns": df.columns.tolist(),
                    "shape": df.shape
                }
            
            return {
                "type": "excel",
                "sheet_names": sheet_names,
                "sheets": sheets_data
            }
            
        except Exception as e:
            logger.error(f"Error processing Excel: {str(e)}", exc_info=True)
            return {"error": f"Failed to process Excel: {str(e)}"}
    
    def _process_powerpoint(self, file_path: str) -> Dict[str, Any]:
        """Process a PowerPoint file."""
        if not PPTX_AVAILABLE:
            return {"error": "python-pptx is not available. Please install with: pip install python-pptx"}
            
        try:
            prs = pptx.Presentation(file_path)
            slides = []
            
            for i, slide in enumerate(prs.slides):
                slide_content = ""
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_content += shape.text + "\n"
                
                slides.append({
                    "slide_number": i + 1,
                    "content": slide_content
                })
            
            return {
                "type": "powerpoint",
                "slide_count": len(slides),
                "slides": slides
            }
            
        except Exception as e:
            logger.error(f"Error processing PowerPoint: {str(e)}", exc_info=True)
            return {"error": f"Failed to process PowerPoint: {str(e)}"}
    
    def _process_archive(self, file_path: str, file_extension: str) -> Dict[str, Any]:
        """Process an archive file."""
        if not ARCHIVE_AVAILABLE:
            return {"error": "zipfile and tarfile modules are required but not available"}
            
        try:
            archive_contents = []
            
            if file_extension == '.zip':
                with zipfile.ZipFile(file_path, 'r') as zip_ref:
                    archive_contents = zip_ref.namelist()
                    
            elif file_extension in ['.tar', '.gz']:
                with tarfile.open(file_path, 'r') as tar_ref:
                    archive_contents = tar_ref.getnames()
            
            return {
                "type": "archive",
                "format": file_extension.lstrip('.'),
                "file_count": len(archive_contents),
                "contents": archive_contents
            }
            
        except Exception as e:
            logger.error(f"Error processing archive: {str(e)}", exc_info=True)
            return {"error": f"Failed to process archive: {str(e)}"}
    
    def _process_unknown_file(self, file_path: str) -> Dict[str, Any]:
        """Process a file with unknown type."""
        try:
            # Try to read as text
            try:
                with open(file_path, 'r', encoding='utf-8') as f:
                    content = f.read()
                return {
                    "type": "text",
                    "content": content
                }
            except UnicodeDecodeError:
                if CHARDET_AVAILABLE:
                    # Try to detect encoding
                    with open(file_path, 'rb') as f:
                        rawdata = f.read()
                    result = chardet.detect(rawdata)
                    encoding = result['encoding']
                    
                    with open(file_path, 'r', encoding=encoding) as f:
                        content = f.read()
                    return {
                        "type": "text",
                        "content": content,
                        "encoding": encoding
                    }
                else:
                    # Return binary file info
                    file_size = os.path.getsize(file_path)
                    mime_type, _ = mimetypes.guess_type(file_path)
                    
                    return {
                        "type": "binary",
                        "size": file_size,
                        "mime_type": mime_type or "application/octet-stream"
                    }
                    
        except Exception as e:
            logger.error(f"Error processing unknown file: {str(e)}", exc_info=True)
            return {"error": f"Failed to process file: {str(e)}"}
    
    def _xml_to_dict(self, element):
        """Convert XML element to dictionary."""
        result = {}
        
        for child in element:
            child_data = self._xml_to_dict(child)
            
            if child.tag in result:
                if type(result[child.tag]) is list:
                    result[child.tag].append(child_data)
                else:
                    result[child.tag] = [result[child.tag], child_data]
            else:
                result[child.tag] = child_data
        
        if element.text and element.text.strip():
            if not result:
                return element.text.strip()
            else:
                result["_text"] = element.text.strip()
        
        if element.attrib:
            if not result:
                result = element.attrib
            else:
                result["_attributes"] = element.attrib
        
        return result